from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

def add_styled_heading(doc, text, level, color=RGBColor(45, 90, 90)):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = color
        run.font.name = 'Arial'

def add_paragraph(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    run.bold = bold
    run.italic = italic
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Arial'

def generate_performance():
    doc = Document()
    doc.add_paragraph('\n' * 5)
    title = doc.add_paragraph('PERFORMANCE ENGINEERING & AUDIT REPORT')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(28)
    title.runs[0].font.bold = True
    title.runs[0].font.color.rgb = RGBColor(45, 90, 90)

    subtitle = doc.add_paragraph('ARTISANT MARKETPLACE PLATFORM')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(18)
    subtitle.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    doc.add_page_break()
    add_styled_heading(doc, '1. Executive Summary', 1)
    add_paragraph(doc, "The Artisant platform is a high-performance marketplace designed for the BTP sector. This report provides a comprehensive analysis of the performance optimizations implemented across the stack.")
    # (Rest of the content...)
    doc.save('deliverables/PerformanceReport_Artisant_TeamArtisant_4TWIN1.docx')

def generate_accessibility():
    doc = Document()
    doc.add_paragraph('\n' * 5)
    title = doc.add_paragraph('DETAILED ACCESSIBILITY AUDIT (WCAG 2.1)')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(28)
    title.runs[0].font.bold = True
    title.runs[0].font.color.rgb = RGBColor(229, 57, 53)

    doc.add_page_break()
    add_styled_heading(doc, '1. Introduction', 1, color=RGBColor(229, 57, 53))
    add_paragraph(doc, "Artisant is committed to ensuring full accessibility for all BTP professionals, including those using assistive technologies.")
    doc.save('deliverables/AccessibilityReport_Artisant_TeamArtisant_4TWIN1.docx')

def generate_ai():
    doc = Document()
    add_styled_heading(doc, 'AI USAGE & INTEGRATION STRATEGY', 0, color=RGBColor(80, 80, 80))
    add_paragraph(doc, "In the development of the Artisant platform, AI was utilized as a force multiplier for complex architectural tasks.")
    doc.save('deliverables/IAUsage_Report_Artisant_TeamArtisant_4TWIN1.docx')

if __name__ == '__main__':
    generate_performance()
    generate_accessibility()
    generate_ai()
    print("Artisant Reports Generated Successfully.")
